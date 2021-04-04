from openpyxl import load_workbook


wb = load_workbook("sampl_random_append1.xlsx")

ws= wb.active

from openpyxl.chart import BarChart,Reference ,LineChart

# bar_value = Reference(ws,min_row=2,max_row=11,min_col=2,max_col=11) ## (ws,범위)
# bar_chart = BarChart()
# bar_chart.add_data(bar_value)  ## 파트 데이터 추가

# ws.add_chart(bar_chart,"E1")

line_value =Reference(ws,min_row=2,max_row=11,min_col=2,max_col=11)
line_chart = LineChart()
line_chart.add_data(line_value,titles_from_data=True) ## 계열> 영어,수학
line_chart.title = "성적표" # 제목
line_chart.style = 20  ## 미리 정의된  스타일
line_chart.y_axis.title = "점수"
line_chart.x_axis.title = "번호"


ws.add_chart(line_chart,"E1")

wb.save("sample_Add_Chart.xlsx")

